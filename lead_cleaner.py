import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import re
import os

def clean_phone(phone):
    if phone == "N/A" or pd.isna(phone):
        return "N/A"
    # Remove all non-numeric characters
    digits = re.sub(r'\D', '', str(phone))
    # Format Indian numbers
    if len(digits) == 10:
        return f"+91 {digits[:5]} {digits[5:]}"
    elif len(digits) == 12 and digits.startswith("91"):
        digits = digits[2:]
        return f"+91 {digits[:5]} {digits[5:]}"
    return str(phone)

def clean_rating(rating):
    if rating == "N/A" or pd.isna(rating):
        return "N/A"
    try:
        # Extract number from rating string
        match = re.search(r'(\d+\.?\d*)', str(rating))
        if match:
            return float(match.group(1))
    except:
        pass
    return "N/A"

def clean_leads(input_file="google_maps_leads.xlsx",
                output_file="cleaned_leads.xlsx"):

    print("📂 Loading raw leads data...")
    df = pd.read_excel(input_file)

    print(f"📊 Raw data: {len(df)} rows")
    print(f"Columns: {list(df.columns)}\n")

    # ── STEP 1: Remove empty rows ──
    print("🧹 Step 1: Removing empty rows...")
    before = len(df)
    df.dropna(how="all", inplace=True)
    print(f"   Removed {before - len(df)} empty rows")

    # ── STEP 2: Remove duplicates ──
    print("🧹 Step 2: Removing duplicate businesses...")
    before = len(df)
    df.drop_duplicates(
        subset=["Business Name"],
        keep="first",
        inplace=True
    )
    print(f"   Removed {before - len(df)} duplicates")

    # ── STEP 3: Clean phone numbers ──
    print("🧹 Step 3: Formatting phone numbers...")
    if "Phone" in df.columns:
        df["Phone"] = df["Phone"].apply(clean_phone)

    # ── STEP 4: Clean ratings ──
    print("🧹 Step 4: Cleaning ratings...")
    if "Rating" in df.columns:
        df["Rating"] = df["Rating"].apply(clean_rating)

    # ── STEP 5: Fill missing values ──
    print("🧹 Step 5: Filling missing values...")
    df.fillna("N/A", inplace=True)
    df.replace("", "N/A", inplace=True)

    # ── STEP 6: Sort by rating ──
    print("🧹 Step 6: Sorting by rating (highest first)...")
    df_rated = df[df["Rating"] != "N/A"].copy()
    df_no_rating = df[df["Rating"] == "N/A"].copy()
    df_rated = df_rated.sort_values(
        by="Rating",
        ascending=False
    )
    df = pd.concat([df_rated, df_no_rating], ignore_index=True)

    # ── STEP 7: Add Lead Quality column ──
    print("🧹 Step 7: Adding lead quality score...")
    def lead_quality(row):
        score = 0
        if row.get("Phone", "N/A") != "N/A":
            score += 1
        if row.get("Website", "N/A") != "N/A":
            score += 1
        if row.get("Rating", "N/A") != "N/A":
            try:
                if float(row["Rating"]) >= 4.0:
                    score += 1
            except:
                pass
        if score == 3:
            return "🔥 Hot"
        elif score == 2:
            return "✅ Good"
        else:
            return "❄️ Cold"

    df["Lead Quality"] = df.apply(lead_quality, axis=1)

    print(f"\n✅ Cleaning complete! {len(df)} clean leads ready\n")

    # ── SAVE TO EXCEL ──
    save_clean_excel(df, output_file)

    return df


def save_clean_excel(df, filename):
    wb = openpyxl.Workbook()

    # ════════════════════════════
    # SHEET 1 — All Cleaned Leads
    # ════════════════════════════
    ws1 = wb.active
    ws1.title = "All Cleaned Leads"

    # Headers
    headers = list(df.columns)
    ws1.append(headers)

    # Header styling
    for cell in ws1[1]:
        cell.font = Font(bold=True, color="FFFFFF", size=11)
        cell.fill = PatternFill("solid", fgColor="1B5E20")
        cell.alignment = Alignment(horizontal="center")

    # Data rows with alternating colors
    for idx, row in df.iterrows():
        ws1.append(list(row))
        row_num = idx + 2
        if idx % 2 == 0:
            for cell in ws1[row_num]:
                cell.fill = PatternFill("solid", fgColor="F1F8E9")

    # Auto width
    for col in ws1.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws1.column_dimensions[
            col[0].column_letter
        ].width = min(max_len + 4, 50)

    # ════════════════════════════
    # SHEET 2 — Hot Leads Only
    # ════════════════════════════
    ws2 = wb.create_sheet("🔥 Hot Leads")

    hot_leads = df[df["Lead Quality"] == "🔥 Hot"]
    ws2.append(headers)

    for cell in ws2[1]:
        cell.font = Font(bold=True, color="FFFFFF", size=11)
        cell.fill = PatternFill("solid", fgColor="B71C1C")
        cell.alignment = Alignment(horizontal="center")

    for idx, row in hot_leads.iterrows():
        ws2.append(list(row))

    for col in ws2.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws2.column_dimensions[
            col[0].column_letter
        ].width = min(max_len + 4, 50)

    # ════════════════════════════
    # SHEET 3 — Summary
    # ════════════════════════════
    ws3 = wb.create_sheet("📊 Summary")

    # Summary data
    total = len(df)
    with_phone = len(df[df["Phone"] != "N/A"])
    without_phone = total - with_phone
    with_website = len(df[df["Website"] != "N/A"])
    hot = len(df[df["Lead Quality"] == "🔥 Hot"])
    good = len(df[df["Lead Quality"] == "✅ Good"])
    cold = len(df[df["Lead Quality"] == "❄️ Cold"])

    rated = df[df["Rating"] != "N/A"]["Rating"]
    try:
        avg_rating = round(float(rated.mean()), 2)
    except:
        avg_rating = "N/A"

    summary_data = [
        ["📊 LEAD GENERATION SUMMARY REPORT", ""],
        ["", ""],
        ["METRIC", "VALUE"],
        ["Total Leads", total],
        ["Leads With Phone Number", with_phone],
        ["Leads Without Phone", without_phone],
        ["Leads With Website", with_website],
        ["Average Rating", avg_rating],
        ["", ""],
        ["LEAD QUALITY BREAKDOWN", ""],
        ["🔥 Hot Leads", hot],
        ["✅ Good Leads", good],
        ["❄️ Cold Leads", cold],
    ]

    for row in summary_data:
        ws3.append(row)

    # Style summary sheet
    ws3["A1"].font = Font(bold=True, size=14, color="1B5E20")
    ws3["A3"].font = Font(bold=True, color="FFFFFF")
    ws3["A3"].fill = PatternFill("solid", fgColor="1B5E20")
    ws3["B3"].font = Font(bold=True, color="FFFFFF")
    ws3["B3"].fill = PatternFill("solid", fgColor="1B5E20")
    ws3["A10"].font = Font(bold=True, color="FFFFFF")
    ws3["A10"].fill = PatternFill("solid", fgColor="1B5E20")

    ws3.column_dimensions["A"].width = 35
    ws3.column_dimensions["B"].width = 20

    wb.save(filename)
    print(f"💾 Saved to {filename}")
    print(f"\n📊 CLEANING SUMMARY:")
    print(f"   Total leads:        {total}")
    print(f"   With phone:         {with_phone}")
    print(f"   Without phone:      {without_phone}")
    print(f"   Hot leads:          {hot}")
    print(f"   Good leads:         {good}")
    print(f"   Average rating:     {avg_rating}")


# ════════ RUN ════════
if __name__ == "__main__":
    print("🚀 Lead Data Cleaning Pipeline")
    print("=" * 40)

    # Input = your Project 1 output
    INPUT_FILE = "google_maps_leads.xlsx"
    OUTPUT_FILE = "cleaned_leads.xlsx"

    if not os.path.exists(INPUT_FILE):
        print(f"❌ File not found: {INPUT_FILE}")
        print("Make sure google_maps_leads.xlsx is in same folder")
    else:
        df = clean_leads(INPUT_FILE, OUTPUT_FILE)
        print(f"\n🎉 Done! Check {OUTPUT_FILE}")